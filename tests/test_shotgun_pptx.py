"""The deck Shotgun writes, opened by something that is not the page.

`tests/shotgun.test.js` builds the .pptx and checks it against the writer's own
idea of what it was writing. That is worth having and it is not enough: a
hand-written OOXML package can be perfectly self-consistent and still be a file
PowerPoint offers to repair. So this reads the same bytes back with `zipfile`
and `xml.etree` - which know nothing about the page - and asks the questions
that actually decide whether it opens:

  * does every part parse as XML,
  * does every relationship resolve to a part that exists,
  * does every `r:embed` on a slide resolve to an image that is really in the
    package, and
  * does the presentation list exactly the slides the package contains.

A dangling `r:embed` is the classic way to produce a "repair" dialog, and it is
invisible until someone double-clicks the file.

python-pptx, if it happens to be importable, is asked to open the file as a
stricter check still. It is deliberately NOT a dependency - the writer is in the
browser and the pipeline should not grow a Python package to test it - so its
absence is a note, not a failure.

Skipped, not failed, when the deck has not been built: it is a product of the
JS suite, and tests/run.sh runs that first.

Run:  bash tests/run.sh          (builds the deck, then runs this)
      python tests/test_shotgun_pptx.py
"""

import os
import sys
import xml.etree.ElementTree as ET
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "build", "shotgun.pptx")

R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
CT_NS = "{http://schemas.openxmlformats.org/package/2006/content-types}"
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def rels_for(part):
    """Path of the .rels part belonging to `part`, by the OPC naming rule."""
    d, name = os.path.split(part)
    return (d + "/" if d else "") + "_rels/" + name + ".rels"


def resolve(base, target):
    """A relationship Target is relative to the part's OWN directory."""
    return os.path.normpath(os.path.join(os.path.dirname(base), target)).replace(os.sep, "/")


def main():
    if not os.path.exists(DECK):
        print("SKIP no tests/build/shotgun.pptx - run tests/run.sh, which builds it")
        return 0

    fails = []
    def chk(label, ok, detail=""):
        print(("ok   " if ok else "FAIL ") + label + (("   " + detail) if detail else ""))
        if not ok:
            fails.append(label)

    with zipfile.ZipFile(DECK) as z:
        names = set(z.namelist())
        # testzip() re-reads every entry and checks its CRC, which is the one
        # thing a hand-written ZIP is most likely to get wrong.
        chk("every entry's CRC is intact", z.testzip() is None)
        chk("stored, not deflated",
            all(i.compress_type == zipfile.ZIP_STORED for i in z.infolist()))

        xml_parts = [n for n in names if n.endswith(".xml") or n.endswith(".rels")]
        bad = []
        for n in xml_parts:
            try:
                ET.fromstring(z.read(n))
            except ET.ParseError as exc:
                bad.append(f"{n}: {exc}")
        chk("every XML part parses", not bad, "; ".join(bad))

        # ---- content types cover every part --------------------------------
        ct = ET.fromstring(z.read("[Content_Types].xml"))
        defaults = {d.get("Extension").lower() for d in ct.findall(CT_NS + "Default")}
        overrides = {o.get("PartName").lstrip("/") for o in ct.findall(CT_NS + "Override")}
        uncovered = [n for n in names
                     if n not in overrides
                     and n.rsplit(".", 1)[-1].lower() not in defaults]
        chk("every part has a content type", not uncovered, ", ".join(uncovered))

        # ---- every relationship resolves -----------------------------------
        dangling, n_rels = [], 0
        for rel_part in [n for n in names if n.endswith(".rels")]:
            owner = rel_part.replace("_rels/", "", 1)[:-5]
            for r in ET.fromstring(z.read(rel_part)).findall(REL_NS + "Relationship"):
                if r.get("TargetMode") == "External":
                    continue
                n_rels += 1
                if resolve(owner, r.get("Target")) not in names:
                    dangling.append(f"{rel_part} -> {r.get('Target')}")
        chk("every relationship target exists", not dangling, ", ".join(dangling))
        print(f"     {n_rels} relationships checked")

        # ---- the slides ----------------------------------------------------
        slides = sorted(n for n in names
                        if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
        pres = ET.fromstring(z.read("ppt/presentation.xml"))
        listed = pres.find(P_NS + "sldIdLst").findall(P_NS + "sldId")
        chk("the presentation lists every slide", len(listed) == len(slides),
            f"{len(listed)} listed, {len(slides)} parts")

        pres_rels = {r.get("Id"): r.get("Target") for r in
                     ET.fromstring(z.read(rels_for("ppt/presentation.xml")))
                     .findall(REL_NS + "Relationship")}
        unlisted = [s.get(R_NS + "id") for s in listed
                    if s.get(R_NS + "id") not in pres_rels]
        chk("every listed slide has a relationship", not unlisted, ", ".join(unlisted))

        # A picture pointing at an r:id the slide's own .rels does not carry is
        # what makes PowerPoint offer to repair the file.
        broken, n_pics = [], 0
        for s in slides:
            rmap = {r.get("Id"): r.get("Target") for r in
                    ET.fromstring(z.read(rels_for(s))).findall(REL_NS + "Relationship")}
            for blip in ET.fromstring(z.read(s)).iter(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
                n_pics += 1
                rid = blip.get(R_NS + "embed")
                if rid not in rmap or resolve(s, rmap[rid]) not in names:
                    broken.append(f"{s}: {rid}")
        chk("every picture on every slide resolves to an image", not broken,
            ", ".join(broken))
        print(f"     {len(slides)} slides, {n_pics} pictures, "
              f"{len([n for n in names if n.startswith('ppt/media/')])} media parts")

        chk("slide size is 16:9",
            pres.find(P_NS + "sldSz").get("cx") == "12192000"
            and pres.find(P_NS + "sldSz").get("cy") == "6858000")

    # ---- the stricter check, when it is available --------------------------
    try:
        from pptx import Presentation
    except ImportError:
        print("note: python-pptx not installed - skipping the library open "
              "(it is not a dependency of this pipeline)")
    else:
        prs = Presentation(DECK)
        chk("python-pptx opens it", len(prs.slides) >= 1)
        shapes = sum(len(s.shapes) for s in prs.slides)
        print(f"     python-pptx: {len(prs.slides)} slides, {shapes} shapes")

    if fails:
        print(f"\n{len(fails)} FAILED")
        return 1
    print("\nALL PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
